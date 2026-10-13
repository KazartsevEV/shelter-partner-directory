"""Read documents actually opened and re-saved by LibreOffice Calc."""
from pathlib import Path
from openpyxl import load_workbook
import math

folder = Path('artifacts/excel')
for name in ['320-ru-linked.xlsx', '320-en-standalone.xlsx', '320-kk-linked.xlsx']:
    source = load_workbook(folder / name)
    converted = load_workbook(folder / 'calc' / name)
    assert source.sheetnames == converted.sheetnames
    for before, after in zip(source, converted):
        assert before.max_row == after.max_row
        for row in before:
            for cell in row:
                other = after[cell.coordinate]
                if isinstance(cell.value, (int, float)):
                    assert isinstance(other.value, (int, float))
                    assert math.isclose(cell.value, other.value, rel_tol=1e-13, abs_tol=1e-12)
                else:
                    assert cell.value == other.value
                assert other.data_type != 'f'
print('EXCEL_LIBREOFFICE_GREEN 3 locales; values, empty cells and literal names preserved')
