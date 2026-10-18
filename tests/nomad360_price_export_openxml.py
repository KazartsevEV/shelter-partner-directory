"""Independent OOXML reading of actual Chromium downloads, not button smoke."""
import json, math, sys, zipfile
from pathlib import Path
from openpyxl import load_workbook

folder = Path(sys.argv[1] if len(sys.argv) > 1 else 'artifacts/excel')
protocol = json.loads((folder / 'protocol.json').read_text())
checked = 0
for case in protocol['cases']:
    for suffix in ['', '-restored']:
        original = Path(case['file'])
        target = folder / (original.stem + suffix + original.suffix)
        book = load_workbook(target, data_only=False)
        assert book.sheetnames == case['sheetNames']
        assert len(book.worksheets) == 2
        sheet = book.worksheets[0]
        assert sheet.max_row == len(case['items']) + 1
        for index, item in enumerate(case['items'], start=2):
            assert sheet.cell(index, 1).value == item['id']
            assert sheet.cell(index, 2).value == item['name']
            assert sheet.cell(index, 4).value == case['currency']
            for column, key in [(5, 'base'), (6, 'discount'), (7, 'customer')]:
                actual = sheet.cell(index, column)
                wanted = item[key]
                if wanted is None:
                    assert actual.value is None
                else:
                    assert actual.data_type == 'n'
                    assert math.isclose(actual.value, wanted, rel_tol=1e-14, abs_tol=1e-12)
            if case['kind'] == 'linked':
                assert sheet.max_column == 12
                for column, key in [(9, 'min'), (10, 'max'), (11, 'target'), (12, 'margin')]:
                    actual = sheet.cell(index, column)
                    wanted = item[key]
                    if wanted is None:
                        assert actual.value is None
                    else:
                        assert actual.data_type == 'n'
                        assert math.isclose(actual.value, wanted, rel_tol=1e-14, abs_tol=1e-12)
            else:
                assert sheet.max_column == 8
            assert sheet.cell(index, 1).data_type == 's'
            assert sheet.cell(index, 2).data_type == 's'
        for tab in book:
            for row in tab:
                for cell in row:
                    assert cell.data_type != 'f'
                    assert cell.hyperlink is None
        with zipfile.ZipFile(target) as archive:
            assert archive.testzip() is None
            assert '[Content_Types].xml' in archive.namelist()
            assert not any('externalLink' in name for name in archive.namelist())
        checked += 1
assert checked == protocol['downloads'] == 72
print('EXCEL_OPENXML_GREEN', json.dumps({'downloads': checked, 'numericCells': True, 'noFormulas': True}))
